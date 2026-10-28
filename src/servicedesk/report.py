"""The monthly SLA report: an Excel workbook for the service desk manager and the monthly service review.

    sdlab report 2026-09 --sample     from the simulated history (the copy in docs/reports)
    sdlab report 2026-09              from the ServiceNow instance

Every number is measured from the tickets' own timestamps against servicenow/config/sla.yaml (see
servicedesk.sla and servicedesk.measure), so the report works the same on the simulated history,
which ServiceNow's SLA engine never saw, and on live tickets. For live tickets the engine's
verdicts are read too and shown beside the report's own, so the two can be checked.

Sheets: Summary, Trend, By priority, Categories, Teams, Departments, Requests, Breaches, Problems
and changes, Tickets (every ticket active in the month, as a filterable table) and Definitions.
"""

from __future__ import annotations

import math
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from servicedesk.business_time import LONDON
from servicedesk.config import Config
from servicedesk.measure import (
    TREND_MONTHS,
    Measured,
    Period,
    add_months,
    describe,
    measure_all,
    median,
    month_bounds,
    percentile_90,
    target_text,
)
from servicedesk.sla import Target, Targets, Ticket
from servicedesk.sources import Extras

# -- the workbook ------------------------------------------------------------------------------

NAVY = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=NAVY)
SUBTITLE_FONT = Font(size=11, color="595959")
SECTION_FONT = Font(bold=True, size=12, color=NAVY)
GOOD = (PatternFill("solid", fgColor="C6EFCE"), Font(color="006100", bold=True))
BAD = (PatternFill("solid", fgColor="FFC7CE"), Font(color="9C0006", bold=True))
FORMATS = {
    "int": "#,##0",
    "pct": "0.0%",
    "target": "0%",
    "hours": "0.0",
    "decimal": "0.0",
    "minutes": "0",
    "datetime": "dd/mm/yyyy hh:mm",
    "date": "dd/mm/yyyy",
    "points": '+0.0" pts";-0.0" pts";0.0" pts"',
}


@dataclass(frozen=True)
class Col:
    header: str
    width: float = 14
    kind: str = "text"  # text, wrap, int, pct, hours, decimal, minutes, datetime, date or points
    align: str | None = None  # left, center or right; numbers sit right and text left by default


def cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.astimezone(LONDON).replace(tzinfo=None)  # Excel has no time zones: show UK time
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def write_table(
    ws: Worksheet, row: int, cols: Sequence[Col], rows: Sequence[Sequence[Any]], *, name: str | None = None
) -> int:
    """A header and its rows from ``row``; an Excel table (filterable) when named. Returns the next free row."""
    lines = 1
    for index, col in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=index, value=col.header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal=col.align)
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = width = max(col.width, ws.column_dimensions[letter].width or 0)
        lines = max(lines, math.ceil(len(col.header) * 1.15 / max(width - 2, 1)))
    ws.row_dimensions[row].height = 15 * max(lines, 2) + 4
    for r, values in enumerate(rows, start=row + 1):
        for index, (col, value) in enumerate(zip(cols, values, strict=True), start=1):
            cell = ws.cell(row=r, column=index, value=cell_value(value))
            if col.kind in FORMATS:
                cell.number_format = FORMATS[col.kind]
            cell.alignment = Alignment(wrap_text=col.kind == "wrap", vertical="top", horizontal=col.align)
    last = row + len(rows)
    if not rows:
        ws.cell(row=row + 1, column=1, value="None this month.").font = SUBTITLE_FONT
        return row + 2
    if name:
        table = Table(displayName=name, ref=f"A{row}:{get_column_letter(len(cols))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
    return last + 1


def mark(cell: Any, met: bool | None) -> None:
    if met is not None:
        cell.fill, cell.font = GOOD if met else BAD


def heading(ws: Worksheet, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


class Report:
    def __init__(self, config: Config, tickets: Sequence[Ticket], extras: Extras, month: datetime) -> None:
        self.config = config
        self.extras = extras
        self.targets = Targets(config)
        self.measured = measure_all(tickets, self.targets, extras.as_at)
        start, _ = month_bounds(month)
        earliest = month_bounds(min((t.opened_at for t in tickets), default=start))[0]
        self.months = [
            Period(self.measured, add_months(start, offset), self.targets)
            for offset in range(-(TREND_MONTHS - 1), 1)
            if add_months(start, offset) >= earliest
        ]
        self.period = self.months[-1]
        self.previous = self.months[-2] if len(self.months) > 1 else None
        categories = config.organisation.get("categories", {})
        self.category_labels = {key: value["label"] for key, value in categories.items()}
        self.subcategory_labels = {
            (key, sub): label
            for key, value in categories.items()
            for sub, label in value["subcategories"].items()
        }
        self.wb = Workbook()
        self.wb.remove(self.wb.worksheets[0])  # start empty: sheets are added in order

    # -- helpers ---------------------------------------------------------------------------------

    def sheet(self, title: str, heading_text: str) -> Worksheet:
        ws = self.wb.create_sheet(title)
        org = self.config.org.name
        heading(
            ws, heading_text, f"{org} IT service desk · {self.period.label} · source: {self.extras.source}"
        )
        return ws

    def working_days(self, target: Target) -> str:
        days = target.seconds / self.targets.working.day_seconds
        return f"{days:g} working day{'s' if days != 1 else ''}"

    def category(self, ticket: Ticket) -> str:
        return self.category_labels.get(ticket.category, ticket.category.title() or "Uncategorised")

    def subcategory(self, ticket: Ticket) -> str:
        return self.subcategory_labels.get((ticket.category, ticket.subcategory), ticket.subcategory)

    def status_at_end(self, m: Measured) -> str:
        ended = m.ended_at
        if ended is None or ended >= self.period.end:
            return "Open"
        t = m.ticket
        if t.state == "cancelled":
            return "Rejected" if t.kind == "request" and t.approval_set else "Cancelled"
        return "Resolved" if t.kind == "incident" else "Completed"

    # -- sheets ----------------------------------------------------------------------------------

    def summary(self) -> str:
        p, prev = self.period, self.previous
        ws = self.sheet("Summary", "IT service desk: monthly SLA report")
        ws["A3"] = (
            f"As at {cell_value(self.extras.as_at):%d/%m/%Y %H:%M}. Every figure is measured from ticket timestamps "
            "against the targets in servicenow/config/sla.yaml: see the Definitions sheet."
        )
        ws["A3"].font = SUBTITLE_FONT
        kpi, previous = p.kpis(), prev.kpis() if prev else None
        rows: list[list[Any]] = []
        rows_met: list[bool | None] = []
        labels = [
            ("response", "Incidents responded to within target"),
            ("resolution", "Incidents resolved within target"),
            ("fulfilment", "Requests completed on time"),
            ("first_contact", "Incidents resolved at first contact"),
            ("automation", "Requests completed by the fulfilment agent"),
        ]
        for key, label in labels:
            rate, target = kpi[key], self.targets.kpi.get(key)
            before = previous[key].value if previous else None
            met = None if target is None or rate.value is None else rate.value * 100 >= target
            change = (rate.value - before) * 100 if rate.value is not None and before is not None else None
            rows.append(
                [
                    label,
                    rate.value,
                    target / 100 if target else None,
                    "" if met is None else ("Met" if met else "Missed"),
                    before,
                    change,
                    f"{rate.met} of {rate.total}" if rate.total else "no tickets",
                ]
            )
            rows_met.append(met)
        counts = [
            ("Incidents logged", len(p.logged("incident")), len(prev.logged("incident")) if prev else None),
            ("Incidents resolved", len(p.resolved()), len(prev.resolved()) if prev else None),
            (
                "Incidents open at month end",
                len(p.open_at_end("incident")),
                len(prev.open_at_end("incident")) if prev else None,
            ),
            (
                "Major incidents (P1)",
                sum(m.ticket.priority == 1 for m in p.logged("incident")),
                sum(m.ticket.priority == 1 for m in prev.logged("incident")) if prev else None,
            ),
            ("Requests logged", len(p.logged("request")), len(prev.logged("request")) if prev else None),
            ("Requests completed", len(p.completed()), len(prev.completed()) if prev else None),
        ]
        cols = [
            Col("Measure", 44),
            Col(p.label, 15, "pct"),
            Col("Target", 10, "target", "center"),
            Col("Status", 10, align="center"),
            Col(prev.label if prev else "Previous month", 15, "pct"),
            Col("Change", 12, "points"),
            Col("Based on", 16, align="center"),
        ]
        top = 5
        next_row = write_table(ws, top, cols, rows)
        for offset, met in enumerate(rows_met, start=1):
            mark(ws.cell(row=top + offset, column=4), met)
        for offset, (label, now_count, before_count) in enumerate(counts):
            r = next_row + offset
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=now_count).number_format = FORMATS["int"]
            if before_count is not None:
                ws.cell(row=r, column=5, value=before_count).number_format = FORMATS["int"]
                ws.cell(row=r, column=6, value=now_count - before_count).number_format = "+#,##0;-#,##0;0"
        notes_row = next_row + len(counts) + 1
        ws.cell(row=notes_row, column=1, value="What happened this month").font = SECTION_FONT
        for offset, note in enumerate(self.notes(), start=1):
            ws.cell(row=notes_row + offset, column=1, value=f"• {note}")
        self.trend_chart(ws, f"I{top}", "Trend")
        return self.headline()

    def headline(self) -> str:
        k = self.period.kpis()
        parts = [
            f"response {k['response']} (target {self.targets.kpi['response']:g}%)",
            f"resolution {k['resolution']} (target {self.targets.kpi['resolution']:g}%)",
            f"requests on time {k['fulfilment']} (target {self.targets.kpi['fulfilment']:g}%)",
        ]
        logged = f"{len(self.period.logged('incident'))} incidents and {len(self.period.logged('request'))} requests logged"
        return f"{self.period.label}: {', '.join(parts)}; {logged}."

    def notes(self) -> list[str]:
        p, prev = self.period, self.previous
        notes = []
        kpi = p.kpis()
        before = prev.kpis() if prev else {}
        last_month = f"{prev.start:%B}" if prev else ""
        for key, label in (
            ("response", "Response"),
            ("resolution", "Resolution"),
            ("fulfilment", "Requests on time"),
        ):
            rate, target = kpi[key], self.targets.kpi[key]
            if rate.value is None:
                continue
            verdict = "met" if rate.value * 100 >= target else "missed"
            change = ""
            earlier = before[key].value if key in before else None
            if earlier is not None:
                points = (rate.value - earlier) * 100
                if abs(points) < 0.05:
                    change = f", level with {last_month}"
                else:
                    change = f", {'up' if points > 0 else 'down'} {abs(points):.1f} points on {last_month}"
            notes.append(f"{label}: {rate.value:.1%} against a {target:g}% target ({verdict}{change}).")

        for m in p.logged("incident"):
            t = m.ticket
            if t.priority != 1:
                continue
            resolution = m.outcome("resolution")
            response = m.outcome("response")
            if resolution and resolution.finished:
                verdict = "within" if resolution.met else "outside"
                opened = t.opened_at.astimezone(LONDON)
                text = (
                    f'P1 {t.number} "{t.short_description}" ({opened.day} {opened:%b}): resolved in '
                    f"{describe(resolution.elapsed)}, {verdict} its {describe(resolution.target.seconds)} target"
                )
            else:
                text = f'P1 {t.number} "{t.short_description}" is still open'
            if response and response.finished and not response.met:
                text += f"; picked up after {describe(response.elapsed)} against {describe(response.target.seconds)}"
            if t.problem:
                text += f"; problem {t.problem}"
            notes.append(text + ".")

        logged = p.logged("incident")
        if logged:
            categories = Counter(self.category(m.ticket) for m in logged)
            top, count = categories.most_common(1)[0]
            change = ""
            if prev and prev.logged("incident"):
                before_count = len(prev.logged("incident"))
                difference = (len(logged) - before_count) / before_count
                change = f" ({'up' if difference >= 0 else 'down'} {abs(difference):.0%} on {last_month})"
            notes.append(
                f"{len(logged)} incidents logged{change}; the most common category was {top} "
                f"({count / len(logged):.0%})."
            )

        breaches = [m for m, o in p.outcomes("resolution") if o.breached]
        if breaches:
            teams = Counter(m.ticket.assignment_group for m in breaches)
            listed = ", ".join(f"{team} {count}" for team, count in teams.most_common())
            notes.append(
                f"{len(breaches)} incidents were resolved outside target: {listed}. Reasons go in the Breaches sheet."
            )

        done, automated = p.completed(), p.automation()
        if done:
            late = [m for m, o in p.outcomes("fulfilment") if o.breached]
            share = f"{automated.met / len(done):.0%}"
            text = f"{len(done)} requests completed, {automated.met} of them ({share}) by the fulfilment agent alone"
            if late:
                text += "; not on time: " + "; ".join(
                    f"{m.ticket.number} {m.ticket.short_description}" for m in late
                )
            rejected = p.rejected()
            if rejected:
                text += f"; {len(rejected)} rejected by the approver"
            notes.append(text + ".")

        new_problems = [x for x in self.extras.problems if self.within(x.opened_at)]
        fixed_problems = [x for x in self.extras.problems if self.within(x.resolved_at)]
        if new_problems or fixed_problems:
            parts = []
            if new_problems:
                parts.append("opened " + ", ".join(f'{x.number} "{x.title}"' for x in new_problems))
            if fixed_problems:
                parts.append("resolved " + ", ".join(x.number for x in fixed_problems))
            notes.append("Problems: " + "; ".join(parts) + ".")
        changes = [c for c in self.extras.changes if self.within(c.start)]
        if changes:
            kinds = Counter(c.type.lower() for c in changes)
            text = f"{len(changes)} changes ({', '.join(f'{n} {k}' for k, n in sorted(kinds.items()))})"
            failed = [c for c in changes if c.outcome == "Unsuccessful"]
            if failed:
                text += "; unsuccessful: " + ", ".join(f'{c.number} "{c.title}"' for c in failed)
            notes.append(text + ".")

        open_late = [m for m, _ in p.open_late() if m.ticket.kind == "incident"]
        open_end = p.open_at_end("incident")
        notes.append(
            f"{len(open_end)} incidents open at month end, {len({m.ticket.number for m in open_late})} of them past target."
        )

        engine = [
            (o.breached, m.ticket.engine_breached[o.target.measure])
            for measure in ("response", "resolution", "fulfilment")
            for m, o in p.outcomes(measure)
            if o.target.measure in m.ticket.engine_breached
        ]
        if engine:
            agree = sum(ours == theirs for ours, theirs in engine)
            notes.append(
                f"ServiceNow's SLA engine agrees with this report on {agree} of {len(engine)} measured SLAs."
            )
        return notes

    def within(self, moment: datetime | None) -> bool:
        return self.period.within(moment)

    def trend(self) -> None:
        ws = self.sheet("Trend", "Trend by month")
        cols = [
            Col("Month", 12),
            Col("Response within target", 13, "pct"),
            Col("Resolution within target", 13, "pct"),
            Col("Requests on time", 13, "pct"),
            Col("First contact resolution", 13, "pct"),
            Col("Completed by the agent", 13, "pct"),
            Col("Incidents logged", 11, "int"),
            Col("Requests logged", 11, "int"),
            Col("Incidents resolved", 11, "int"),
            Col("Requests completed", 11, "int"),
            Col("Open incidents at month end", 13, "int"),
        ]
        rows = []
        for month in self.months:
            k = month.kpis()
            rows.append(
                [
                    f"{month.start:%b %Y}",
                    k["response"].value,
                    k["resolution"].value,
                    k["fulfilment"].value,
                    k["first_contact"].value,
                    k["automation"].value,
                    len(month.logged("incident")),
                    len(month.logged("request")),
                    len(month.resolved()),
                    len(month.completed()),
                    len(month.open_at_end("incident")),
                ]
            )
        last = write_table(ws, 4, cols, rows)
        ws.freeze_panes = "B5"
        targets = self.targets.kpi
        ws.cell(
            row=last + 1,
            column=1,
            value=(
                f"Targets: response {targets['response']:g}%, resolution {targets['resolution']:g}%, "
                f"requests on time {targets['fulfilment']:g}%."
            ),
        ).font = SUBTITLE_FONT
        self.trend_chart(ws, f"A{last + 3}", "Trend")
        volumes = BarChart()
        volumes.type = "col"
        volumes.title = "Tickets logged by month"
        volumes.y_axis.title = "Tickets"
        volumes.add_data(
            Reference(ws, min_col=7, max_col=8, min_row=4, max_row=4 + len(rows)), titles_from_data=True
        )
        volumes.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(rows)))
        volumes.width, volumes.height = 18, 8
        ws.add_chart(volumes, f"H{last + 3}")

    def trend_chart(self, ws: Worksheet, anchor: str, data_sheet: str) -> None:
        source = self.wb[data_sheet] if data_sheet in self.wb.sheetnames else None
        if source is None:
            return
        rows = len(self.months)
        chart = LineChart()
        chart.title = "Share within target, by month"
        chart.y_axis.number_format = "0%"
        chart.y_axis.scaling.max = 1.0
        lowest = min(
            (
                v
                for month in self.months
                for k, rate in month.kpis().items()
                if k in ("response", "resolution", "fulfilment") and (v := rate.value) is not None
            ),
            default=0.8,
        )
        chart.y_axis.scaling.min = max(0.0, (int(lowest * 10) - 1) / 10)
        chart.add_data(
            Reference(source, min_col=2, max_col=4, min_row=4, max_row=4 + rows), titles_from_data=True
        )
        chart.set_categories(Reference(source, min_col=1, min_row=5, max_row=4 + rows))
        chart.width, chart.height = 18, 8
        for series in chart.series:
            series.smooth = False
        ws.add_chart(chart, anchor)

    def by_priority(self) -> None:
        ws = self.sheet("By priority", "Incidents by priority")
        p = self.period
        cols = [
            Col("Priority", 14),
            Col("Response target", 16),
            Col("Resolution target", 16),
            Col("Logged", 9, "int"),
            Col("Resolved", 9, "int"),
            Col("Response within target", 12, "pct"),
            Col("Resolution within target", 12, "pct"),
            Col("Median response (min)", 12, "minutes"),
            Col("Median time to resolve (h)", 12, "hours"),
            Col("90% resolved within (h)", 12, "hours"),
            Col("Response breaches", 10, "int"),
            Col("Resolution breaches", 10, "int"),
        ]
        rows = []
        for priority in sorted({*self.targets.labels, 5}):
            label = f"P{priority} {self.targets.labels.get(priority, 'Planning')}"
            targets = self.targets.incident.get(priority)

            def same(m: Measured, priority: int = priority) -> bool:
                return m.ticket.priority == priority

            logged = [m for m in p.logged("incident") if same(m)]
            resolved = [m for m in p.resolved() if same(m)]
            if not targets:
                rows.append(
                    [
                        label,
                        "No SLA",
                        "No SLA",
                        len(logged),
                        len(resolved),
                        None,
                        None,
                        None,
                        None,
                        None,
                        None,
                        None,
                    ]
                )
                continue
            response, resolution = p.outcomes("response", same), p.outcomes("resolution", same)
            hours = [o.elapsed / 3600 for _, o in resolution]
            rows.append(
                [
                    label,
                    target_text(targets["response"]),
                    target_text(targets["resolution"]),
                    len(logged),
                    len(resolved),
                    p.rate("response", same).value,
                    p.rate("resolution", same).value,
                    median([o.elapsed / 60 for _, o in response]),
                    median(hours),
                    percentile_90(hours),
                    sum(o.breached for _, o in response),
                    sum(o.breached for _, o in resolution),
                ]
            )
        write_table(ws, 4, cols, rows)
        ws.cell(
            row=5 + len(rows) + 1,
            column=1,
            value=(
                "Times are on each target's clock: working hours (08:00-18:00 on working days) for P2 to P4, "
                "all hours for P1."
            ),
        ).font = SUBTITLE_FONT

    def categories(self) -> None:
        ws = self.sheet("Categories", "Incidents by category")
        p = self.period
        logged = p.logged("incident")
        keys = sorted(
            {(m.ticket.category, m.ticket.subcategory) for m in logged + p.resolved()},
            key=lambda key: (-sum((m.ticket.category, m.ticket.subcategory) == key for m in logged), key),
        )
        cols = [
            Col("Category", 20),
            Col("Subcategory", 26),
            Col("Logged", 9, "int"),
            Col("Share of incidents", 11, "pct"),
            Col("Resolved", 9, "int"),
            Col("Resolved at first contact", 12, "pct"),
            Col("Median time to resolve (h)", 12, "hours"),
            Col("Resolution breaches", 11, "int"),
        ]
        rows = []
        for category, subcategory in keys:

            def same(m: Measured, key: tuple[str, str] = (category, subcategory)) -> bool:
                return (m.ticket.category, m.ticket.subcategory) == key

            count = sum(same(m) for m in logged)
            resolved = [m for m in p.resolved() if same(m)]
            first = sum(self.targets.first_contact(m.ticket) for m in resolved)
            outcomes = p.outcomes("resolution", same)
            sample = next(m.ticket for m in logged + p.resolved() if same(m))
            rows.append(
                [
                    self.category(sample),
                    self.subcategory(sample),
                    count,
                    count / len(logged) if logged else None,
                    len(resolved),
                    first / len(resolved) if resolved else None,
                    median([o.elapsed / 3600 for _, o in outcomes]),
                    sum(o.breached for _, o in outcomes),
                ]
            )
        write_table(ws, 4, cols, rows, name="Categories")
        ws.freeze_panes = "A5"

    def teams(self) -> None:
        ws = self.sheet("Teams", "Incidents by resolving team")
        p = self.period
        groups = sorted({m.ticket.assignment_group for m in p.resolved() + p.open_at_end("incident")} - {""})
        cols = [
            Col("Team", 24),
            Col("Resolved", 10, "int"),
            Col("Resolution within target", 12, "pct"),
            Col("Median time to resolve (h)", 12, "hours"),
            Col("Passed on at least once", 12, "pct"),
            Col("Open at month end", 10, "int"),
            Col("Open and past target", 10, "int"),
        ]
        open_late = {m.ticket.number for m, o in p.open_late() if m.ticket.kind == "incident"}
        rows = []
        for group in groups:

            def same(m: Measured, group: str = group) -> bool:
                return m.ticket.assignment_group == group

            resolved = [m for m in p.resolved() if same(m)]
            outcomes = p.outcomes("resolution", same)
            open_end = [m for m in p.open_at_end("incident") if same(m)]
            rows.append(
                [
                    group,
                    len(resolved),
                    p.rate("resolution", same).value,
                    median([o.elapsed / 3600 for _, o in outcomes]),
                    sum(m.ticket.reassignments > 0 for m in resolved) / len(resolved) if resolved else None,
                    len(open_end),
                    sum(m.ticket.number in open_late for m in open_end),
                ]
            )
        write_table(ws, 4, cols, rows, name="Teams")

    def departments(self) -> None:
        ws = self.sheet("Departments", "Demand by department")
        p = self.period
        staff = Counter(
            person.department
            for person in self.config.org.people
            if person.start_date < p.end.date()
            and (person.end_date is None or person.end_date >= p.start.date())
        )
        incidents, requests = p.logged("incident"), p.logged("request")
        names = sorted(set(staff) | {m.ticket.department for m in incidents + requests} - {""})
        cols = [
            Col("Department", 26),
            Col("Staff", 8, "int"),
            Col("Incidents logged", 11, "int"),
            Col("Incidents per 10 staff", 11, "decimal"),
            Col("Requests logged", 11, "int"),
            Col("Most common incident category", 28),
        ]
        rows = []
        for name in names:
            mine = [m for m in incidents if m.ticket.department == name]
            common = Counter(self.category(m.ticket) for m in mine).most_common(1)
            rows.append(
                [
                    name,
                    staff.get(name, 0),
                    len(mine),
                    len(mine) / staff[name] * 10 if staff.get(name) else None,
                    sum(m.ticket.department == name for m in requests),
                    common[0][0] if common else "",
                ]
            )
        write_table(ws, 4, cols, rows, name="Departments")

    def requests(self) -> None:
        ws = self.sheet("Requests", "Service requests by catalogue item")
        p = self.period
        items = [i["name"] for i in self.config.catalog["items"]]
        items += sorted({m.ticket.item for m in p.active() if m.ticket.kind == "request"} - set(items) - {""})
        cols = [
            Col("Catalogue item", 28),
            Col("Logged", 9, "int"),
            Col("Rejected", 9, "int"),
            Col("Completed", 10, "int"),
            Col("Completed by the agent", 11, "int"),
            Col("On time", 9, "int"),
            Col("On time %", 10, "pct"),
            Col("Median wait for approval (h)", 13, "hours"),
            Col("Median time to complete after approval (h)", 15, "hours"),
            Col("Measured as", 30, "wrap"),
        ]
        rows = []
        clock = self.targets.working
        no_approval = {i["name"] for i in self.config.catalog["items"] if i.get("approval") == "none"}
        for item in items:

            def same(m: Measured, item: str = item) -> bool:
                return m.ticket.item == item

            done = [m for m in p.completed() if same(m)]
            outcomes = p.outcomes("fulfilment", same)
            waits = [
                clock.seconds_between(m.ticket.opened_at, m.ticket.approval_set) / 3600
                for m in done
                if m.ticket.approval_set
            ]
            dated = item in self.targets.dated_items
            work = [o.elapsed / 3600 for _, o in outcomes if not o.target.on_date]
            rows.append(
                [
                    item,
                    sum(same(m) for m in p.logged("request")),
                    sum(same(m) for m in p.rejected()),
                    len(done),
                    sum(m.ticket.automated for m in done),
                    sum(o.met for _, o in outcomes),
                    p.rate("fulfilment", same).value,
                    None if item in no_approval else median(waits),
                    None if dated else median(work),
                    "Done by the end of the due date"
                    if dated
                    else f"{self.working_days(self.targets.fulfilment)} from approval",
                ]
            )
        write_table(ws, 4, cols, rows, name="Requests")

    def breaches(self) -> None:
        ws = self.sheet("Breaches", "SLA breaches for the service review")
        p = self.period
        cols = [
            Col("Ticket", 14),
            Col("Type", 9),
            Col("Priority", 8),
            Col("Measure", 11),
            Col("Target", 18),
            Col("Due", 16, "datetime"),
            Col("Finished", 16, "datetime"),
            Col("Over by (h)", 9, "hours"),
            Col("Status", 16),
            Col("Team", 22),
            Col("Assigned to", 16),
            Col("Category", 18),
            Col("Short description", 50, "wrap"),
            Col("Reason (for the service review)", 36, "wrap"),
        ]
        rows = []
        for m, o in sorted(p.breaches(), key=lambda pair: pair[1].due):
            t = m.ticket
            rows.append(
                [
                    t.number,
                    t.kind.title(),
                    f"P{t.priority}" if t.kind == "incident" else "",
                    o.target.measure.title(),
                    target_text(o.target, o.due),
                    o.due,
                    o.stopped,
                    o.over_seconds / 3600,
                    "Breached",
                    t.assignment_group,
                    t.assigned_to,
                    self.category(t) if t.kind == "incident" else t.item,
                    t.short_description,
                    "",
                ]
            )
        for m, o in sorted(p.open_late(), key=lambda pair: pair[1].due):
            t = m.ticket
            over = o.target.clock.seconds_between(o.due, p.end)
            rows.append(
                [
                    t.number,
                    t.kind.title(),
                    f"P{t.priority}" if t.kind == "incident" else "",
                    o.target.measure.title(),
                    target_text(o.target, o.due),
                    o.due,
                    None,
                    over / 3600,
                    "Open, past target",
                    t.assignment_group,
                    t.assigned_to,
                    self.category(t) if t.kind == "incident" else t.item,
                    t.short_description,
                    "",
                ]
            )
        write_table(ws, 4, cols, rows, name="Breaches")
        ws.freeze_panes = "B5"

    def problems_and_changes(self) -> None:
        ws = self.sheet("Problems and changes", "Problems and changes")
        p = self.period
        linked_month = Counter(m.ticket.problem for m in p.logged("incident") if m.ticket.problem)
        linked_all = Counter(m.ticket.problem for m in self.measured if m.ticket.problem)
        problems = [
            x
            for x in self.extras.problems
            if x.opened_at < p.end and (x.resolved_at is None or x.resolved_at >= p.start)
        ]
        ws.cell(row=4, column=1, value="Problems open during the month").font = SECTION_FONT
        cols = [
            Col("Problem", 14),
            Col("Title", 44, "wrap"),
            Col("State", 16),
            Col("Team", 20),
            Col("Opened", 16, "datetime"),
            Col("Resolved", 16, "datetime"),
            Col("Known error", 10),
            Col("Incidents this month", 11, "int"),
            Col("Incidents in total", 11, "int"),
        ]
        rows = [
            [
                x.number,
                x.title,
                x.state,
                x.group,
                x.opened_at,
                x.resolved_at,
                x.known_error,
                linked_month.get(x.number, 0),
                linked_all.get(x.number, 0),
            ]
            for x in problems
        ]
        row = write_table(ws, 5, cols, rows, name="Problems")
        changes = [c for c in self.extras.changes if self.within(c.start) or self.within(c.end)]
        ws.cell(row=row + 1, column=1, value="Changes carried out in the month").font = SECTION_FONT
        cols = [
            Col("Change", 14),
            Col("Type", 12),
            Col("Title", 44, "wrap"),
            Col("Planned start", 16, "datetime"),
            Col("Planned end", 16, "datetime"),
            Col("State", 12),
            Col("Risk", 10),
            Col("Outcome", 14),
        ]
        rows = [[c.number, c.type, c.title, c.start, c.end, c.state, c.risk, c.outcome] for c in changes]
        write_table(ws, row + 2, cols, rows, name="Changes")

    def ticket_list(self) -> None:
        ws = self.sheet("Tickets", "Every ticket active in the month")
        p = self.period
        cols = [
            Col("Ticket", 14),
            Col("Type", 9),
            Col("Opened", 16, "datetime"),
            Col("Priority", 8),
            Col("Category", 18),
            Col("Subcategory", 22),
            Col("Catalogue item", 22),
            Col("Team", 22),
            Col("Assigned to", 16),
            Col("Department", 22),
            Col("Channel", 12),
            Col("Status at month end", 12),
            Col("Response (min)", 10, "minutes"),
            Col("Response met", 9),
            Col("Resolved or completed", 16, "datetime"),
            Col("Time to resolve (h)", 10, "hours"),
            Col("Within target", 9),
            Col("First contact", 9),
            Col("By the agent", 9),
            Col("Problem", 14),
            Col("SLA engine agrees", 10),
            Col("Short description", 60),
        ]
        rows = []
        for m in p.active():
            t = m.ticket
            response = m.outcome("response")
            finish = m.outcome("resolution") or m.outcome("fulfilment")
            engine = [
                m.ticket.engine_breached[o.target.measure] == o.breached
                for o in m.outcomes
                if o.finished and o.target.measure in m.ticket.engine_breached
            ]
            incident = t.kind == "incident"
            rows.append(
                [
                    t.number,
                    t.kind.title(),
                    t.opened_at,
                    f"P{t.priority}" if incident else "",
                    self.category(t) if incident else "",
                    self.subcategory(t) if incident else "",
                    t.item,
                    t.assignment_group,
                    t.assigned_to,
                    t.department,
                    t.contact_type,
                    self.status_at_end(m),
                    response.elapsed / 60 if response and response.finished else None,
                    response.met if response and response.finished else None,
                    m.ended_at if m.ended_at and m.ended_at < p.end else None,
                    finish.elapsed / 3600
                    if finish and finish.finished and not finish.target.on_date
                    else None,
                    finish.met if finish and finish.finished else None,
                    self.targets.first_contact(t) if incident else None,
                    t.automated if not incident else None,
                    t.problem,
                    all(engine) if engine else None,
                    t.short_description,
                ]
            )
        write_table(ws, 4, cols, rows, name="Tickets")
        ws.freeze_panes = "B5"

    def definitions(self) -> None:
        ws = self.sheet("Definitions", "How each figure is measured")
        t = self.targets
        schedule = self.config.schedule
        rows = [
            (
                "Month",
                "Each event counts in the calendar month (UK time) in which it happened: logged, picked up, "
                "resolved or completed. A ticket logged on 30 September and resolved on 2 October counts as logged in "
                "September and resolved in October.",
            ),
            (
                "Working hours",
                f"{schedule['start']} to {schedule['end']}, Monday to Friday, except bank holidays in "
                'England and Wales (ServiceNow schedule "Kestrel service desk hours"). P1 targets run all hours.',
            ),
            (
                "Response",
                "From logging until an analyst starts work (Actual start, recorded when the ticket first goes "
                "to In Progress) or resolves it. Targets: "
                + "; ".join(f"P{p} {target_text(v['response'])}" for p, v in sorted(t.incident.items()))
                + ".",
            ),
            (
                "Resolution",
                "From logging until Resolved, less any time On Hold (waiting for the caller or a supplier). "
                "Targets: "
                + "; ".join(f"P{p} {target_text(v['resolution'])}" for p, v in sorted(t.incident.items()))
                + ". P5 (Planning) has no SLA: the date is agreed with the caller.",
            ),
            (
                "Request fulfilment",
                f"From approval (or logging, when none is needed) until the request is completed: "
                f"{self.working_days(t.fulfilment)}. Waiting for the approver does not count against the desk.",
            ),
            (
                "Done on the due date",
                "Leavers and moves have a date of their own. They are on time if completed by the "
                "end of that day (a leaver's account disabled on their last day), however early they were logged.",
            ),
            (
                "First contact resolution",
                f"Resolved by the {t.first_contact_group} itself, never passed to another "
                f"team, within {t.first_contact_within.total_seconds() / 60:.0f} minutes of being logged.",
            ),
            (
                "Completed by the agent",
                "Completed requests that the fulfilment agent on the domain controller carried "
                "out in Active Directory with no one touching them. The rest were handed to a person.",
            ),
            (
                "Breach",
                "A target missed. Breaches count in the month the ticket was resolved or completed. Tickets "
                'still open at month end that were already past target are listed as "Open, past target" and count '
                "when they finish.",
            ),
            (
                "Not measured",
                "Cancelled incidents and rejected requests (nothing was owed); tickets raised for people "
                "outside the council's staff list; tickets created by the automated smoke test.",
            ),
            (
                "SLA engine agrees",
                'For live tickets, ServiceNow\'s own SLA records are read as well. "Yes" means its '
                "verdict matches this report's. Simulated history (SIM- correlation IDs) is not timed by the engine.",
            ),
            (
                "Source",
                f"{self.extras.source.capitalize()}, as at {cell_value(self.extras.as_at):%d/%m/%Y %H:%M}. "
                "Made by `sdlab report` from the it-service-desk-lab repository.",
            ),
        ]
        write_table(ws, 4, [Col("Term", 26, "wrap"), Col("Definition", 110, "wrap")], [list(r) for r in rows])

    def build(self, path: Path) -> str:
        self.trend()  # first: the summary's chart reads the trend table
        headline = self.summary()
        self.wb.move_sheet("Summary", offset=-1)
        self.by_priority()
        self.categories()
        self.teams()
        self.departments()
        self.requests()
        self.breaches()
        self.problems_and_changes()
        self.ticket_list()
        self.definitions()
        self.wb.active = 0
        props = self.wb.properties
        props.title = f"IT service desk SLA report, {self.period.label}"
        props.creator = "sdlab report"
        props.created = props.modified = self.extras.as_at.astimezone(UTC).replace(tzinfo=None)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)
        return headline


def build_report(
    config: Config, tickets: Sequence[Ticket], extras: Extras, *, month: datetime, path: Path
) -> str:
    """Writes the workbook and returns a one-line headline."""
    return Report(config, tickets, extras, month).build(path)
